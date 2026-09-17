import os
import time

import pandas as pd
import openpyxl
from openpyxl import load_workbook

DetailedReport = "JioSignage_FocusedTest.xlsx"


def convertCSVtoXLSX(csv_file, xlsx_file):
    try:
        cvsDataframe = pd.read_csv(csv_file)
        resultExcelFile = pd.ExcelWriter(xlsx_file)
        cvsDataframe.to_excel(resultExcelFile, index=False)
        resultExcelFile.close()
    except Exception as e:
        print(e)


def delete_columns_by_names(file_path, sheet_name, column_names):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]
    for column_name in column_names:
        column_index = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == column_name:
                column_index = col
                break
        if column_index is not None:
            ws.delete_cols(column_index)
        else:
            pass
    wb.save(file_path)


def updateResultInTestReport(sheet_name):
    # Load the workbooks
    workbook1 = openpyxl.load_workbook('xl_detailed_test.xlsx')
    workbook2 = openpyxl.load_workbook(DetailedReport)

    # Choose the sheets
    sheet1 = workbook1['Sheet1']
    sheet2 = workbook2[sheet_name]

    # Get the maximum row number for each sheet
    max_row_sheet1 = sheet1.max_row
    max_row_sheet2 = sheet2.max_row

    for row_num_sheet1 in range(1, max_row_sheet1 + 1):
        value_sheet1 = sheet1.cell(row=row_num_sheet1, column=3).value

        for row_num_sheet2 in range(1, max_row_sheet2 + 1):
            value_sheet2 = sheet2.cell(row=row_num_sheet2, column=8).value
            if value_sheet1 == value_sheet2:
                value_to_copy = sheet1.cell(row=row_num_sheet1, column=6).value

                sheet2.cell(row=row_num_sheet2, column=9).value = value_to_copy
                break  # Break the inner loop once a match is found

    workbook2.save(DetailedReport)
    print(sheet_name + " is updated successfully.")


if __name__ == '__main__':
    columns = ["Start Time", "Stop Time", "Duration in ms", "Parent Suite", "Suite", "Test Class", "Test Method",
               "Description", "Sub Suite"]
    convertCSVtoXLSX('csv_detailed_test.csv', 'xl_detailed_test.xlsx')
    delete_columns_by_names('xl_detailed_test.xlsx', "Sheet1", columns)
    time.sleep(2)
    updateResultInTestReport("Login & Profile")
    updateResultInTestReport("Dashboard")
    updateResultInTestReport("Materials")
    updateResultInTestReport("Layouts")
    updateResultInTestReport("Playlists")
    updateResultInTestReport("Schedules")
    updateResultInTestReport("Displays")
    updateResultInTestReport("User Access")
    updateResultInTestReport("Display Status")
    updateResultInTestReport("Emergenecy Alerts")

